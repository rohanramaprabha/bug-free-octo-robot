"""
PrivacyVault AI - DOCX Document Extractor
Extracts paragraphs, table cells, and document properties from DOCX files using python-docx.
"""

from pathlib import Path
import docx
from .extractor_base import ExtractedDocument, DocumentSegment


def extract_docx(file_path: Path, filename: str) -> ExtractedDocument:
    """Extract paragraphs, tables, and metadata properties from a Word document (.docx)."""
    doc = ExtractedDocument(filename=filename, file_type="docx")
    try:
        word_doc = docx.Document(str(file_path))
        
        # 1. Metadata extraction from core properties
        props = word_doc.core_properties
        meta_dict = {}
        if props.author:
            meta_dict["author"] = props.author
        if props.last_modified_by:
            meta_dict["last_modified_by"] = props.last_modified_by
        if props.title:
            meta_dict["title"] = props.title
        if props.revision:
            meta_dict["revision"] = str(props.revision)
        if props.created:
            meta_dict["created"] = str(props.created)
        if props.comments:
            meta_dict["comments"] = props.comments
            
        doc.metadata = meta_dict
        
        segments = []
        full_text_parts = []
        
        # 2. Extract paragraphs
        for p_idx, p in enumerate(word_doc.paragraphs, start=1):
            text = p.text.strip()
            if text:
                segments.append(DocumentSegment(
                    text=text,
                    location_type="section",
                    location_ref=f"Paragraph {p_idx}",
                    meta={"paragraph": p_idx}
                ))
                full_text_parts.append(text)
                
        # 3. Extract tables
        for t_idx, table in enumerate(word_doc.tables, start=1):
            headers = []
            for r_idx, row in enumerate(table.rows, start=1):
                cell_texts = [cell.text.strip() for cell in row.cells]
                if r_idx == 1:
                    headers = cell_texts
                
                for c_idx, cell_val in enumerate(cell_texts, start=1):
                    if cell_val:
                        col_name = headers[c_idx - 1] if c_idx <= len(headers) else f"Col {c_idx}"
                        segments.append(DocumentSegment(
                            text=cell_val,
                            location_type="row_col",
                            location_ref=f"Table {t_idx}, Row {r_idx}, Col '{col_name}'",
                            meta={"table": t_idx, "row": r_idx, "col": col_name}
                        ))
                        full_text_parts.append(cell_val)
                        
        doc.segments = segments
        doc.raw_text = "\n".join(full_text_parts)
        doc.metadata["total_paragraphs"] = len(word_doc.paragraphs)
        doc.metadata["total_tables"] = len(word_doc.tables)
        doc.is_empty = len(segments) == 0
        
    except Exception as e:
        doc.error = f"DOCX extraction failure: {str(e)}"
        
    return doc
